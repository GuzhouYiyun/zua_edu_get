import requests, re, time, os, sys
from hashlib import sha1

os.system("chcp 65001 >nul")

url = "http://jwglxt.zua.edu.cn/eams/courseTableForStd!courseTable.action"
login_url = "http://jwglxt.zua.edu.cn/eams/loginExt.action"

# 网页实例
s = requests.Session()

#----------登录页面
page = s.get(login_url).text

username = input("学号: ")
password = input("密码: ")

#密钥
m = re.search(r"SHA1\('([0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}-)'", page)
sha1p = sha1((m.group(1) + password).encode()).hexdigest()

#----------登录
headers = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/151.0.0.0 Safari/537.36 Edg/151.0.0.0",
    "Referer": "http://jwglxt.zua.edu.cn/eams/loginExt.action",
    "Origin": "http://jwglxt.zua.edu.cn",
    "Content-Type": "application/x-www-form-urlencoded",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
    "Accept-Language": "zh-CN,zh;q=0.9",
    "Upgrade-Insecure-Requests": "1",
    }

data = {
    "needModify": "0",
    "username": username,
    "password": sha1p,
    "session_locale": "zh_CN",
    "login_from": "login_from"
    }

s.post(login_url, headers=headers, data=data)

# 检查登录是否成功
r = s.get("http://jwglxt.zua.edu.cn/eams/courseTableForStd.action")
if "semesterBar" not in r.text:
    print("登录失败！请检查学号和密码。")
    os.system("pause")
    sys.exit()
else:
    print("登录成功！请稍等...")

#查询时间id
tag_id = "semesterBar" + re.search(r'semesterBar(\d+)Semester', r.text).group(1) + "Semester"
cur = s.cookies.get("semester.id") or "297"
r2 = s.post(url = "http://jwglxt.zua.edu.cn/eams/dataQuery.action", data = {
    "tagId": tag_id,
    "dataType": "semesterCalendar",
    "value": cur,
    "empty": "false"
    })
beg = input("输入学期年份后两位（如2026-2027的输入26即可）：")
sx = input("输入学期（1上学期，2下学期）：")
tcode = re.search(rf'id:(\d+),schoolYear:"{"20" + beg + "-20" + str(int(beg) + 1)}",name:"{sx}"', r2.text).group(1)

headers = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/151.0.0.0 Safari/537.36 Edg/151.0.0.0",
    "Referer": "http://jwglxt.zua.edu.cn/eams/courseTableForStd.action",
    "Origin": "http://jwglxt.zua.edu.cn",
    "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
    "Accept": "*/*",
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8,en-GB;q=0.7,en-US;q=0.6",
    "x-requested-with": "XMLHttpRequest"
    }

ids = re.search(r'"ids","(\d+)"', r.text).group(1)

data = {
    "ignoreHead": "1",
    "setting.kind": "std",
    "startWeek": "",
    "semester.id": tcode,
    "ids": ids
    }

#获取课表
resp = s.post(url = url, headers = headers, data = data)

#print(resp.text)

# ---------- 解析课表并生成 Excel（追加部分） ----------
import openpyxl

def split_args(s):
    out, buf, depth, q = [], "", 0, None
    for c in s:
        if q:
            buf += c
            if c == q:
                q = None
        elif c in '"\'':
            buf += c
            q = c
        elif c == '(':
            depth += 1
            buf += c
        elif c == ')':
            depth -= 1
            buf += c
        elif c == ',' and depth == 0:
            out.append(buf.strip())
            buf = ""
        else:
            buf += c
    if buf.strip():
        out.append(buf.strip())
    return out

ta_pat = re.compile(r"new TaskActivity\((.*?)\);", re.S)
idx_pat = re.compile(r"index\s*=\s*(\d+)\s*\*\s*unitCount\s*\+\s*(\d+)\s*;")
grid = [[None] * 7 for _ in range(10)]  # grid[节次][星期几]

ta_positions = [m.start() for m in ta_pat.finditer(resp.text)]

for i, ta_pos in enumerate(ta_positions):
    ta_match = ta_pat.search(resp.text, ta_pos)
    args = split_args(ta_match.group(1))
    if len(args) < 7:
        continue
    course = args[3].strip('"').split("(")[0] or args[2].strip('"').split("(")[0]
    room = args[5].strip('"') or args[4].strip('"')
    bits = args[6].strip('"')
    weeks = [str(j) for j, ch in enumerate(bits) if ch == "1"]

    next_ta = ta_positions[i + 1] if i + 1 < len(ta_positions) else len(resp.text)
    index_block = resp.text[ta_match.end():next_ta]
    units = idx_pat.findall(index_block)

    seg = resp.text[:ta_pos]
    tpos = [t.start() for t in re.finditer(r"var teachers = \[", seg)]
    if tpos:
        names = re.findall(r'name:"([^"]+)"', seg[tpos[-1]:])
        teacher = "、".join(dict.fromkeys(names)) or "未知"
    else:
        teacher = "未知"

    ent = (course, room, teacher, ",".join(weeks))
    if units:
        day = int(units[0][0])
        unit_vals = sorted(set(int(u) for _, u in units))
        for u in unit_vals:
            if u < 10 and day < 7:
                if grid[u][day] is None:
                    grid[u][day] = [ent]
                else:
                    for j, e in enumerate(grid[u][day]):
                        if e[0] == course and e[1] == room and e[2] == teacher:
                            merged = sorted(set(e[3].split(",")) | set(weeks), key=int)
                            grid[u][day][j] = (e[0], e[1], e[2], ",".join(merged))
                            break
                    else:
                        grid[u][day].append(ent)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "课表"
days = ["星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日"]
al = openpyxl.styles.Alignment(horizontal="center", vertical="center", wrap_text=True)

ws.cell(1, 1, "节次/周次")
for d, txt in enumerate(days):
    ws.cell(1, 2 + d, txt)

merged = set()

for d in range(7):
    u = 0
    while u < 10:
        if (u, d) in merged or grid[u][d] is None:
            u += 1
            continue

        entries = grid[u][d]
        entry_set = frozenset((e[0], e[1], e[2]) for e in entries)

        group = [u]
        for next_u in range(u + 1, 10):
            if (next_u, d) in merged:
                break
            if grid[next_u][d] is None:
                break
            next_set = frozenset((e[0], e[1], e[2]) for e in grid[next_u][d])
            if entry_set != next_set:
                break
            prev = group[-1]
            if (prev == 3 and next_u == 4) or (prev == 7 and next_u == 8):
                break
            group.append(next_u)

        lines = []
        for course, room, teacher, w in entries:
            lines += [course, room, teacher, w]
        val = "\n".join(lines)
        r, c = 2 + group[0], 2 + d
        if len(group) > 1:
            end_row = 2 + group[-1]
            ws.merge_cells(start_row=r, start_column=c, end_row=end_row, end_column=c)
            for mu in group[1:]:
                merged.add((mu, d))
        ws.cell(r, c, val).alignment = al
        u = group[-1] + 1

for u in range(10):
    ws.cell(2 + u, 1, "第" + str(u + 1) + "节").alignment = al
for i, wd in enumerate([10] + [17] * 7, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = wd
for u in range(10):
    ws.row_dimensions[2 + u].height = 72

folder = os.path.join(os.path.expanduser("~"), "Desktop", "课表")
os.makedirs(folder, exist_ok=True)

out = os.path.join(folder, "课表.xlsx")
wb.save(out)
print("已生成Excel课表：", out)

# 生成TXT课表（按TaskActivity块合并连续节次）
txt_lines = []
days = ["星期一", "星期二", "星期三", "星期四", "星期五", "星期六", "星期日"]

for i, ta_pos in enumerate(ta_positions):
    ta_match = ta_pat.search(resp.text, ta_pos)
    args = split_args(ta_match.group(1))
    if len(args) < 7:
        continue
    course = args[3].strip('"').split("(")[0] or args[2].strip('"').split("(")[0]
    room = args[5].strip('"') or args[4].strip('"')
    bits = args[6].strip('"')
    weeks = [str(j) for j, ch in enumerate(bits) if ch == "1"]

    next_ta = ta_positions[i + 1] if i + 1 < len(ta_positions) else len(resp.text)
    index_block = resp.text[ta_match.end():next_ta]
    units = idx_pat.findall(index_block)

    seg = resp.text[:ta_pos]
    tpos = [t.start() for t in re.finditer(r"var teachers = \[", seg)]
    if tpos:
        names = re.findall(r'name:"([^"]+)"', seg[tpos[-1]:])
        teacher = "、".join(dict.fromkeys(names)) or "未知"
    else:
        teacher = "未知"

    if not units:
        continue

    day_periods = {}
    for day_str, slot_str in units:
        day = int(day_str)
        slot = int(slot_str)
        if day not in day_periods:
            day_periods[day] = []
        day_periods[day].append(slot)

    for day, slots in day_periods.items():
        if day >= 7:
            continue
        slots = sorted(set(slots))
        groups = []
        current_group = [slots[0]]
        for j in range(1, len(slots)):
            prev = slots[j - 1]
            curr = slots[j]
            if curr == prev + 1 and not (prev == 3 and curr == 4) and not (prev == 7 and curr == 8):
                current_group.append(curr)
            else:
                groups.append(current_group)
                current_group = [curr]
        groups.append(current_group)

        for group in groups:
            if len(group) == 1:
                time_str = f"时间：{days[day]} 第{group[0] + 1}节"
            else:
                time_str = f"时间：{days[day]} 第{group[0] + 1}-{group[-1] + 1}节"
            txt_lines.append(course)
            txt_lines.append(time_str)
            txt_lines.append(f"地点：{room}")
            txt_lines.append(f"教师：{teacher}")
            txt_lines.append(f"周次：{','.join(weeks)}")
            txt_lines.append("")

txt_out = os.path.join(folder, "课表.txt")
with open(txt_out, "w", encoding="utf-8") as f:
    f.write("\n".join(txt_lines))
print("已生成TXT课表：", txt_out)
os.system("pause")